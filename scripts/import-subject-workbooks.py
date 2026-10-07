"""Extract the six official 2569/2 catalogs. Usage: python script.py INPUT_DIRECTORY."""
import json
import re
import sys
from pathlib import Path

import openpyxl


def text(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def extract(path):
    sheet = openpyxl.load_workbook(path, data_only=True).active
    # Resolve only explicitly merged cells; do not carry values across sections.
    cells = {(r, c): sheet.cell(r, c).value for r in range(1, sheet.max_row + 1)
             for c in range(1, sheet.max_column + 1)}
    for merged in sheet.merged_cells.ranges:
        for r in range(merged.min_row, merged.max_row + 1):
            for c in range(merged.min_col, merged.max_col + 1):
                cells[r, c] = cells[merged.min_row, merged.min_col]
    headers = {text(cells[2, c]): c for c in range(1, sheet.max_column + 1)}
    seat_col = headers['จำนวนรับ']
    condition_col = headers['เงื่อนไขรายวิชา']
    time_cols = list(range(condition_col + 1, sheet.max_column + 1))
    subjects = []
    expected_marks = 0
    for r in range(5, sheet.max_row + 1):
        expected_marks += sum(text(sheet.cell(r, c).value) == '✔' for c in time_cols)
        code = text(cells[r, 1])
        if not re.fullmatch(r'[ก-๙A-Za-z]+\d{5}', code):
            assert not any(text(sheet.cell(r, c).value) == '✔' for c in time_cols), (path, r)
            continue
        slots, alternate, restrictions = [], [], []
        for c in time_cols:
            if text(sheet.cell(r, c).value) != '✔':
                continue
            day = text(cells[2, c])
            period = text(cells[3, c])
            match = re.fullmatch(r'(\d+)(?:\s*\(([^)]+)\))?', period)
            assert match, (path, r, c, period)
            slots.append(f'{day} {match[1]}')
            if match[2]:
                alternate.append(f'{day} {match[1]} → {match[2]}')
            restriction = text(cells[4, c])
            if restriction:
                restrictions.append(f'{day} {match[1]}: {restriction}')
        assert slots, (path, r, code, 'no periods')
        notes = []
        condition = text(cells[r, condition_col])
        if condition and condition != '-':
            notes.append(condition)
        if 'กลุ่มการเรียนที่ลงทะเบียนเรียนได้' in headers:
            notes.append('กลุ่มการเรียนที่ลงทะเบียนเรียนได้: ' + text(cells[r, headers['กลุ่มการเรียนที่ลงทะเบียนเรียนได้']]))
        notes.extend(restrictions)
        if alternate:
            notes.append('ตารางคาบ 0, 9 (4 พ.ย. 69 - 11 ธ.ค. 69): ' + ', '.join(alternate))
        notes.append('ห้ามลงทะเบียนเรียนซ้ำในรายวิชาที่เคยเรียนมาแล้วในช่วง ' + ('ม.1-3' if sheet.title in ['ม.1', 'ม.2', 'ม.3'] else 'ม.4-6'))
        seats = text(cells[r, seat_col])
        subjects.append(dict(
            order=r, code=code, name=re.sub(r'\s+', ' ', text(cells[r, 2])),
            credit=text(cells[r, 3]), classPerWeek=text(cells[r, 4]),
            description=text(cells[r, 5]), instructor=text(cells[r, 6]),
            group=text(cells[r, 7]), enrollment=seats, electiveQuantity=seats,
            updatedElectiveQuantity=seats, classtime=', '.join(slots),
            classroom='', note='\n'.join(notes),
        ))
    assert sum(len(s['classtime'].split(', ')) for s in subjects) == expected_marks
    identities = [(s['code'], s['group'], s['classtime']) for s in subjects]
    assert len(identities) == len(set(identities)), (path, 'duplicate groups')
    return subjects


if __name__ == '__main__':
    source_dir = Path(sys.argv[1])
    root = Path(__file__).resolve().parents[1]
    catalog = {}
    for grade in range(1, 7):
        paths = list(source_dir.glob(f'ม.{grade} *2569.xlsx'))
        assert len(paths) == 1, (grade, paths)
        catalog[str(grade)] = extract(paths[0])
        print(f'ม.{grade}: {len(catalog[str(grade)])} course groups, '
              f'{len(set(s["code"] for s in catalog[str(grade)]))} course codes')
    payload = json.dumps(catalog, ensure_ascii=False, indent=2) + '\n'
    for target in ['src/data/subjects-2569-2.json', 'server/src/subjects-2569-2.json']:
        destination = root / target
        destination.parent.mkdir(parents=True, exist_ok=True)
        destination.write_text(payload)
